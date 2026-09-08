from io import BytesIO

MAX_RESUME_BYTES = 5 * 1024 * 1024
MAX_RESUME_CHARS = 12000


class ResumeParseError(ValueError):
    pass


def extract_resume_text(filename: str, data: bytes) -> str:
    if not data:
        raise ResumeParseError("The resume file is empty.")
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeParseError("Resume is too large. Please upload a file under 5 MB.")

    name = (filename or "").lower().strip()
    try:
        if name.endswith(".pdf"):
            text = _from_pdf(data)
        elif name.endswith(".docx"):
            text = _from_docx(data)
        elif name.endswith(".txt"):
            text = data.decode("utf-8", errors="replace")
        else:
            raise ResumeParseError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("Could not read this resume. The file may be corrupted.") from exc

    cleaned = "\n".join(line.strip() for line in text.splitlines()).strip()
    if len(cleaned) < 20:
        raise ResumeParseError("No readable text was found in this resume.")
    return cleaned[:MAX_RESUME_CHARS]


def _from_pdf(data: bytes) -> str:
    import fitz

    document = fitz.open(stream=data, filetype="pdf")
    try:
        return "\n".join(page.get_text() or "" for page in document)
    finally:
        document.close()


def _from_docx(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(parts)
